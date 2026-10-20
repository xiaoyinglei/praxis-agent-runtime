from __future__ import annotations

import os as _os_module
import tempfile
from collections.abc import Sequence
from contextlib import AbstractContextManager, nullcontext
from dataclasses import dataclass, field
from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path
from typing import Any, Protocol, cast

import duckdb as _duckdb

from rag.ingest.asset_anchors import asset_anchor, iter_asset_anchor_refs
from rag.ingest.parsers.dispatcher import ExtractionDispatcher
from rag.ingest.parsers.util import normalize_whitespace
from rag.ingest.parsers.web_parser_repo import WebParserRepo
from rag.ingest.retrievalsummarizer import RetrievalSummarizer, RetrievalSummaryResult
from rag.ingest.section_refiner import SectionRefiner
from rag.ingest.table_sampler import TableAssetProfile, profile_markdown_table
from rag.schema.core import (
    AssetRecord,
    AssetSummaryRecord,
    DocSummaryRecord,
    Document,
    DocumentStatus,
    LayoutMetaCacheRecord,
    ParsedDocument,
    ParsedElement,
    ParsedSection,
    PartitionKey,
    ProcessingStateRecord,
    SectionLocatorRecord,
    SectionRecord,
    SectionSummaryRecord,
    Source,
    SourceType,
    StorageTier,
)
from rag.schema.model_protocols import Embedder
from rag.schema.runtime import AccessPolicy, ObjectStore


def _normalize_visible_text_for_locator(text: str) -> str:
    """
    只做保守标准化，确保：
    1. locator 计算
    2. visible_text 落盘
    3. runtime grounding 回读
    三者使用同一份文本基底。
    """
    if not text:
        return ""

    # 去 BOM
    text = text.replace("\ufeff", "")

    # 统一换行
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # 去掉少量脏字符，但不做激进空白折叠
    text = text.replace("\x00", "")

    return text
def _parse_markdown_table_to_dataframe(markdown: str) -> Any:
    """Parse a markdown table string into a pandas DataFrame with all rows."""
    import pandas as pd

    normalized = markdown.replace("\r\n", "\n").replace("\r", "\n").strip()
    rows: list[list[str]] = []
    for line in normalized.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if all(set(cell) <= {"-", ":", " "} for cell in cells if cell):
            continue
        rows.append(cells)
    if len(rows) < 2:
        return None
    header, data_rows = rows[0], rows[1:]
    max_cols = max(len(header), max((len(r) for r in data_rows), default=0))
    columns = [col or f"col_{i}" for i, col in enumerate(header)]
    while len(columns) < max_cols:
        columns.append(f"col_{len(columns)}")
    padded_rows = [r + [""] * (max_cols - len(r)) for r in data_rows]
    return pd.DataFrame(padded_rows, columns=columns)


# ============================================================
# Pipeline-facing narrow protocols
# ============================================================

class MetadataWriterRepo(Protocol):
    def save_source(self, source: Source) -> Source: ...
    def save_document(self, document: Document) -> Document: ...
    def save_section(self, section: SectionRecord) -> SectionRecord: ...
    def save_asset(self, asset: AssetRecord) -> AssetRecord: ...
    def save_layout_meta_cache(self, record: LayoutMetaCacheRecord) -> LayoutMetaCacheRecord: ...
    def save_processing_state(self, record: ProcessingStateRecord) -> ProcessingStateRecord: ...
    def set_document_index_state(
        self,
        doc_id: int,
        *,
        is_indexed: bool | None = None,
        index_ready: bool | None = None,
        indexed_at: datetime | None = None,
        last_index_error: str | None = None,
    ) -> Document: ...


class SummaryIndexRepo(Protocol):
    """
    L2 summary + vector writer.

    约定：
    - 主键稳定（section_id / asset_id / doc_id）
    - 普通重试依赖 upsert 覆盖
    """
    def upsert_record(
        self,
        record: DocSummaryRecord | SectionSummaryRecord | AssetSummaryRecord,
        vector: Sequence[float],
        *,
        embedding_space: str = "default",
    ) -> None: ...

    def upsert_records(
        self,
        items: Sequence[tuple[DocSummaryRecord | SectionSummaryRecord | AssetSummaryRecord, Sequence[float]]],
        *,
        embedding_space: str = "default",
    ) -> None: ...


# ============================================================
# Request / Result
# ============================================================

@dataclass(frozen=True, slots=True)
class IngestRequest:
    location: str
    source_type: SourceType | str
    owner: str
    title: str | None = None
    content_text: str | None = None
    raw_bytes: bytes | None = None
    file_path: Path | None = None
    metadata: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class DirectContentItem:
    location: str
    source_type: SourceType | str
    content: str | bytes | Path
    owner: str = "user"
    title: str | None = None
    metadata: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class IngestPipelineResult:
    source_id: int
    doc_id: int
    title: str
    status: str
    section_count: int
    asset_count: int

    @property
    def indexed_object_count(self) -> int:
        return self.section_count + self.asset_count + 1


@dataclass(frozen=True, slots=True)
class BatchIngestItemResult:
    request: IngestRequest
    result: IngestPipelineResult | None = None
    error: str | None = None

    @property
    def succeeded(self) -> bool:
        return self.result is not None and self.error is None

    @property
    def indexed_object_count(self) -> int:
        if self.result is None:
            return 0
        return self.result.indexed_object_count


@dataclass(frozen=True, slots=True)
class BatchIngestResult:
    results: list[BatchIngestItemResult]

    @property
    def success_count(self) -> int:
        return sum(1 for item in self.results if item.succeeded)

    @property
    def failure_count(self) -> int:
        return sum(1 for item in self.results if not item.succeeded)

    @property
    def indexed_object_count(self) -> int:
        return sum(item.indexed_object_count for item in self.results)


@dataclass(frozen=True, slots=True)
class _PreparedIngestItem:
    request: IngestRequest
    source: Source
    document: Document
    parsed_doc: ParsedDocument
    saved_sections: list[SectionRecord]
    saved_assets: list[AssetRecord]
# ============================================================
# Pipeline
# ============================================================

class IngestPipeline:
    """
    两阶段工业级 ingest pipeline

    Phase 1: L1 事实落库
        Source / Document / SectionRecord / AssetRecord -> PostgreSQL

    Phase 2: L2 索引构建
        summary -> embedding -> summary+vector upsert -> Milvus

    幂等策略：
    - 普通重试：依赖稳定主键 + upsert 覆盖
    - 重建任务：由外部专门 worker 先清理再全量重建
    """

    def __init__(
        self,
        dispatcher: ExtractionDispatcher,
        summarizer: RetrievalSummarizer,
        embedder: Embedder,
        metadata_repo: MetadataWriterRepo,
        summary_repo: SummaryIndexRepo,
        object_store: ObjectStore | None = None,
        embedding_model_id: str = "default",
        embedding_space: str = "default",
        section_refiner: SectionRefiner | None = None,
    ) -> None:
        self._dispatcher = dispatcher
        self._summarizer = summarizer
        self._embedder = embedder
        self._metadata_repo = metadata_repo
        self._summary_repo = summary_repo
        self._object_store = object_store
        self._embedding_model_id = embedding_model_id or "default"
        self._embedding_space = embedding_space or "default"
        self._section_refiner = section_refiner or SectionRefiner()

    def configure_summarizer(self, summarizer: RetrievalSummarizer) -> None:
        self._summarizer = summarizer

    def run(self, request: IngestRequest) -> IngestPipelineResult:
        with self._metadata_transaction():
            prepared = self._prepare_l1(request)
        self._write_l2_batch([prepared])
        return self._result_from_prepared(prepared)

    def run_many(self, requests: Sequence[IngestRequest]) -> list[IngestPipelineResult]:
        if not requests:
            return []
        with self._metadata_transaction():
            prepared_items = [self._prepare_l1(request) for request in requests]
        self._write_l2_batch(prepared_items)
        return [self._result_from_prepared(prepared) for prepared in prepared_items]

    def _prepare_l1(self, request: IngestRequest) -> _PreparedIngestItem:
        source_type = SourceType(request.source_type)
        raw_bytes = self._resolve_raw_bytes(request)
        content_hash = sha256(raw_bytes).hexdigest()

        # -------------------------
        # Phase 1: L1 事实层落库
        # -------------------------
        source = self._build_source(
            request=request,
            source_type=source_type,
            content_hash=content_hash,
            raw_bytes=raw_bytes,
        )
        source = self._metadata_repo.save_source(source)

        try:
            parsed_doc = self._parse_document(
                request=request,
                source_type=source_type,
                raw_bytes=raw_bytes,
            )
        except Exception as exc:
            # At this point document does not exist, so there is no processing_state to persist.
            raise ValueError(f"parse failed for {request.location}: {exc}") from exc

        parsed_doc = self._section_refiner.refine(parsed_doc)
        asset_kinds_by_ref = self._asset_kinds_by_ref(parsed_doc.elements)

        raw_object_key = None
        visible_text_key = None

        visible_text = parsed_doc.visible_text
        if not visible_text:
            raise ValueError(
                f"parsed_doc.visible_text is empty for location={request.location}"
            )

        if self._object_store is not None:
            suffix = Path(request.location).suffix

            # 1) 原始文件单独存储
            raw_object_key = self._object_store.put_bytes(raw_bytes, suffix=suffix)

            # 2) visible_text 单独存储
            visible_text_key = self._object_store.put_bytes(
                visible_text.encode("utf-8"),
                suffix=".visible.txt",
            )
            source = self._metadata_repo.save_source(source.model_copy(update={"object_key": raw_object_key}))

        document = self._build_document(
            source=source,
            parsed_doc=parsed_doc,
            content_hash=content_hash,
        )
        document = self._metadata_repo.save_document(document)

        if raw_object_key is None:
            raise ValueError(
                f"raw object storage key is missing for location={request.location}"
            )

        if visible_text_key is None:
            raise ValueError(
                f"visible_text storage key is missing for location={request.location}"
            )

        section_locators = self._build_section_locators(
            visible_text=visible_text,
            sections=parsed_doc.sections,
            visible_text_key=visible_text_key,
        )

        saved_sections: list[SectionRecord] = []
        for order_index, (parsed_section, locator) in enumerate(
            zip(parsed_doc.sections, section_locators, strict=True)
        ):
            if parsed_section.order_index != order_index:
                raise ValueError(
                    f"section order mismatch: parsed_section.order_index={parsed_section.order_index}, "
                    f"loop order_index={order_index}, toc_path={parsed_section.toc_path!r}"
                )
            section_record = self._build_section_record(
                source=source,
                document=document,
                parsed_section=parsed_section,
                order_index=order_index,
                content_storage_key=raw_object_key,
                locator=locator,
                asset_kinds_by_ref=asset_kinds_by_ref,
            )
            saved_sections.append(self._metadata_repo.save_section(section_record))

        saved_sections = self._bind_refined_section_parent_ids(saved_sections)

        section_id_by_anchor_ref = self._section_ids_by_anchor_ref(
            parsed_sections=parsed_doc.sections,
            saved_sections=saved_sections,
        )

        saved_assets: list[AssetRecord] = []
        for parsed_element in parsed_doc.elements:
            asset_record = self._build_asset_record(
                source=source,
                document=document,
                parsed_element=parsed_element,
                saved_sections=saved_sections,
                section_id_by_anchor_ref=section_id_by_anchor_ref,
                content_storage_key=raw_object_key,
            )
            if asset_record is None:
                continue
            saved_assets.append(self._metadata_repo.save_asset(asset_record))

        layout_cache = self._build_layout_meta_cache_record(
            source=source,
            document=document,
            parsed_doc=parsed_doc,
            content_hash=content_hash,
            object_key=raw_object_key,
        )
        if layout_cache is not None:
            self._metadata_repo.save_layout_meta_cache(layout_cache)

        # L1 安全点已经完成，进入 indexing
        self._save_processing_state(
            doc_id=document.doc_id,
            source_id=source.source_id,
            stage="index",
            status="processing",
        )

        return _PreparedIngestItem(
            request=request,
            source=source,
            document=document,
            parsed_doc=parsed_doc,
            saved_sections=saved_sections,
            saved_assets=saved_assets,
        )

    def _write_l2_batch(self, prepared_items: Sequence[_PreparedIngestItem]) -> None:
        if not prepared_items:
            return

        try:
            self._write_prepared_summary_records(prepared_items)

            indexed_at = datetime.now(UTC)
            with self._metadata_transaction():
                for prepared in prepared_items:
                    self._mark_l2_ready(prepared, indexed_at=indexed_at)

        except Exception as exc:
            with self._metadata_transaction():
                for prepared in prepared_items:
                    self._mark_l2_failed(prepared, error_message=str(exc))
            raise

    def _metadata_transaction(self) -> AbstractContextManager[Any]:
        transaction = getattr(self._metadata_repo, "transaction", None)
        if callable(transaction):
            return cast("AbstractContextManager[Any]", transaction())
        return nullcontext()

    def _mark_l2_ready(self, prepared: _PreparedIngestItem, *, indexed_at: datetime) -> None:
        self._metadata_repo.set_document_index_state(
            prepared.document.doc_id,
            is_indexed=True,
            index_ready=True,
            indexed_at=indexed_at,
            last_index_error=None,
        )

        self._save_processing_state(
            doc_id=prepared.document.doc_id,
            source_id=prepared.source.source_id,
            stage="index",
            status="ready",
        )

    def _mark_l2_failed(self, prepared: _PreparedIngestItem, *, error_message: str) -> None:
        self._metadata_repo.set_document_index_state(
            prepared.document.doc_id,
            is_indexed=False,
            index_ready=False,
            indexed_at=None,
            last_index_error=error_message,
        )
        self._save_processing_state(
            doc_id=prepared.document.doc_id,
            source_id=prepared.source.source_id,
            stage="index",
            status="failed",
            error_message=error_message,
        )

    @staticmethod
    def _result_from_prepared(prepared: _PreparedIngestItem) -> IngestPipelineResult:
        return IngestPipelineResult(
            source_id=prepared.source.source_id,
            doc_id=prepared.document.doc_id,
            title=prepared.parsed_doc.title,
            status="ready",
            section_count=len(prepared.saved_sections),
            asset_count=len(prepared.saved_assets),
        )

    # ============================================================
    # L1 builders
    # ============================================================

    def _parse_document(
        self,
        *,
        request: IngestRequest,
        source_type: SourceType,
        raw_bytes: bytes,
    ) -> ParsedDocument:
        if source_type is SourceType.WEB and request.content_text is not None:
            return WebParserRepo().parse(
                raw_bytes.decode("utf-8", errors="replace"),
                location=request.location,
                title=request.title,
                owner=request.owner,
            )
        if source_type in {
            SourceType.PLAIN_TEXT,
            SourceType.PASTED_TEXT,
            SourceType.BROWSER_CLIP,
        }:
            return self._parse_plain_text(
                request=request,
                source_type=source_type,
                raw_bytes=raw_bytes,
            )
        return self._dispatcher.route_and_parse(
            file_path=request.file_path or Path(request.location),
            location=request.location,
            source_type=source_type,
            title=request.title,
            owner=request.owner,
        )

    def _parse_plain_text(
        self,
        *,
        request: IngestRequest,
        source_type: SourceType,
        raw_bytes: bytes,
    ) -> ParsedDocument:
        title = request.title or Path(request.location).name or request.location
        visible_text = _normalize_visible_text_for_locator(raw_bytes.decode("utf-8", errors="replace")).strip() or title
        metadata = {
            **{str(key): str(value) for key, value in request.metadata.items()},
            "location": request.location,
            "source_type": source_type.value,
        }
        section = ParsedSection(
            toc_path=(title,),
            heading_level=1,
            page_range=None,
            order_index=0,
            text=visible_text,
            char_range_start=0,
            char_range_end=len(visible_text),
            anchor_hint=None,
            metadata=metadata,
        )
        return ParsedDocument(
            title=title,
            source_type=source_type,
            authors=[request.owner],
            language=None,
            sections=[section],
            visible_text=visible_text,
            visual_semantics=None,
            elements=[],
            page_count=None,
            metadata=metadata,
        )

    def _build_source(
        self,
        *,
        request: IngestRequest,
        source_type: SourceType,
        content_hash: str,
        raw_bytes: bytes,
    ) -> Source:
        return Source(
            source_type=source_type,
            location=request.location,
            original_file_name=Path(request.location).name or None,
            content_hash=content_hash,
            file_size_bytes=len(raw_bytes),
            owner_id=request.owner,
            effective_access_policy=AccessPolicy.default(),
            metadata_json={str(key): str(value) for key, value in request.metadata.items()},
        )

    def _build_document(
        self,
        *,
        source: Source,
        parsed_doc: ParsedDocument,
        content_hash: str,
    ) -> Document:
        return Document(
            source_id=source.source_id,
            title=parsed_doc.title,
            language=parsed_doc.language,
            authors=list(parsed_doc.authors),
            file_hash=content_hash,
            version_group_id=0,
            version_no=1,
            doc_status=DocumentStatus.PUBLISHED,
            effective_date=None,
            is_active=True,
            is_indexed=False,
            index_ready=False,
            index_priority="high",
            storage_tier=StorageTier.HOT,
            reference_count=1,
            page_count=parsed_doc.page_count,
            tenant_id=None,
            department_id=None,
            auth_tag=None,
            embedding_model_id=self._embedding_model_id,
            indexed_at=None,
            last_index_error=None,
            effective_access_policy=AccessPolicy.default(),
            metadata_json=dict(parsed_doc.metadata),
        )
    def _build_section_locators(
        self,
        *,
        visible_text: str,
        sections: Sequence[ParsedSection],
        visible_text_key: str,
    ) -> list[SectionLocatorRecord]:
        """
        基于 parser 已经产出的 char span 构建 section locator。

        正式版要求：
        1. parser 必须提供完整合法的 char_range_start / char_range_end
        2. pipeline 不再做文本查找
        3. 若 span 非法或 text/span 不一致，直接报错
        """
        if not visible_text_key:
            raise ValueError("visible_text_key must not be empty")

        locators: list[SectionLocatorRecord] = []

        for idx, section in enumerate(sections):
            char_start = section.char_range_start
            char_end = section.char_range_end

            if char_start is None or char_end is None:
                raise ValueError(
                    f"section[{idx}] missing char range: "
                    f"toc_path={section.toc_path!r}, order_index={section.order_index}"
                )

            if char_start < 0:
                raise ValueError(
                    f"section[{idx}] invalid char_range_start={char_start}: "
                    f"toc_path={section.toc_path!r}, order_index={section.order_index}"
                )

            if char_end <= char_start:
                raise ValueError(
                    f"section[{idx}] invalid char range start={char_start}, end={char_end}: "
                    f"toc_path={section.toc_path!r}, order_index={section.order_index}"
                )

            if char_end > len(visible_text):
                raise ValueError(
                    f"section[{idx}] char_range_end out of bounds: "
                    f"end={char_end}, visible_text_len={len(visible_text)}, "
                    f"toc_path={section.toc_path!r}, order_index={section.order_index}"
                )

            sliced_text = visible_text[char_start:char_end]
            if sliced_text != section.text:
                raise ValueError(
                    f"section[{idx}] text/span mismatch: "
                    f"toc_path={section.toc_path!r}, order_index={section.order_index}, "
                    f"expected={section.text!r}, actual={sliced_text!r}"
                )

            byte_start = len(visible_text[:char_start].encode("utf-8"))
            byte_end = len(visible_text[:char_end].encode("utf-8"))

            locators.append(
                SectionLocatorRecord(
                    visible_text_key=visible_text_key,
                    char_range_start=char_start,
                    char_range_end=char_end,
                    byte_range_start=byte_start,
                    byte_range_end=byte_end,
                )
            )

        return locators

    def _build_section_record(
        self,
        *,
        source: Source,
        document: Document,
        parsed_section: ParsedSection,
        order_index: int,
        content_storage_key: str | None,
        locator: SectionLocatorRecord,
        asset_kinds_by_ref: dict[str, str],
    ) -> SectionRecord:
        page_start = parsed_section.page_range[0] if parsed_section.page_range else None
        page_end = parsed_section.page_range[1] if parsed_section.page_range else None
        if not locator.visible_text_key:
            raise ValueError("locator.visible_text_key must not be empty")

        if locator.char_range_start < 0:
            raise ValueError("locator.char_range_start must be >= 0")

        if locator.char_range_end <= locator.char_range_start:
            raise ValueError("locator char range is invalid")

        if locator.byte_range_start < 0:
            raise ValueError("locator.byte_range_start must be >= 0")

        if locator.byte_range_end < locator.byte_range_start:
            raise ValueError("locator byte range is invalid")
        section_text = parsed_section.text
        content_hash = sha256(section_text.encode("utf-8")).hexdigest()
        anchor_refs = iter_asset_anchor_refs(section_text)
        anchor_kinds = [
            asset_kinds_by_ref[anchor_ref]
            for anchor_ref in anchor_refs
            if anchor_ref in asset_kinds_by_ref
        ]
        has_table = any(kind == "table" for kind in anchor_kinds)
        has_figure = any(kind in {"figure", "image_summary"} for kind in anchor_kinds)
        metadata_json: dict[str, Any] = dict(parsed_section.metadata)
        metadata_json.setdefault("document_title", document.title or "")
        metadata_json.setdefault(
            "section_title",
            parsed_section.toc_path[-1] if parsed_section.toc_path else document.title or "",
        )
        refined_window_index = self._metadata_int(metadata_json, "refined_window_index")
        refined_window_count = self._metadata_int(metadata_json, "refined_window_count")
        metadata_json["window_index"] = refined_window_index if refined_window_index is not None else 0
        metadata_json["window_count"] = refined_window_count if refined_window_count is not None else 1
        if anchor_refs:
            metadata_json["asset_anchor_refs"] = list(anchor_refs)

        return SectionRecord(
            doc_id=document.doc_id,
            source_id=source.source_id,
            parent_section_id=None,
            toc_path=list(parsed_section.toc_path),
            heading_level=parsed_section.heading_level,
            order_index=order_index,
            anchor=parsed_section.anchor_hint,
            page_start=page_start,
            page_end=page_end,
            content_storage_key=content_storage_key,
            visible_text_key=locator.visible_text_key,
            raw_locator=locator,
            char_range_start=locator.char_range_start,
            char_range_end=locator.char_range_end,
            byte_range_start=locator.byte_range_start,
            byte_range_end=locator.byte_range_end,
            section_kind="section",
            content_hash=content_hash,
            has_table=has_table,
            has_figure=has_figure,
            neighbor_asset_count=len(anchor_kinds),
            metadata_json=metadata_json,
        )

    def _bind_refined_section_parent_ids(self, saved_sections: list[SectionRecord]) -> list[SectionRecord]:
        groups: dict[tuple[int, int, tuple[str, ...], int], list[SectionRecord]] = {}
        for section in saved_sections:
            group_key = self._refined_section_group_key(section)
            if group_key is None:
                continue
            groups.setdefault(group_key, []).append(section)

        if not groups:
            return saved_sections

        updated_by_id: dict[int, SectionRecord] = {}
        for sections in groups.values():
            if len(sections) <= 1:
                continue
            ordered = sorted(
                sections,
                key=lambda section: (
                    self._section_window_index(section) if self._section_window_index(section) is not None else 0,
                    section.order_index,
                    section.section_id,
                ),
            )
            parent_section_id = ordered[0].section_id
            window_count = self._section_window_count(ordered)
            for fallback_index, section in enumerate(ordered):
                window_index = self._section_window_index(section)
                metadata_json = {
                    **section.metadata_json,
                    "logical_section_id": parent_section_id,
                    "parent_section_id": parent_section_id,
                    "window_index": window_index if window_index is not None else fallback_index,
                    "window_count": window_count,
                    "neighbor_expansion_boundary": "parent_section_id",
                }
                updated_section = section.model_copy(
                    update={
                        "parent_section_id": parent_section_id,
                        "metadata_json": metadata_json,
                    }
                )
                updated_by_id[section.section_id] = self._metadata_repo.save_section(updated_section)

        if not updated_by_id:
            return saved_sections
        return [updated_by_id.get(section.section_id, section) for section in saved_sections]

    def _refined_section_group_key(
        self,
        section: SectionRecord,
    ) -> tuple[int, int, tuple[str, ...], int] | None:
        metadata = section.metadata_json
        if str(metadata.get("refine_strategy", "") or "").strip() != "token_window":
            return None
        refined_from = self._metadata_int(metadata, "refined_from_section_order")
        if refined_from is None:
            return None
        return (section.doc_id, section.source_id, tuple(section.toc_path), refined_from)

    def _section_window_index(self, section: SectionRecord) -> int | None:
        for key in ("window_index", "refined_window_index"):
            value = self._metadata_int(section.metadata_json, key)
            if value is not None:
                return value
        return None

    def _section_window_count(self, sections: Sequence[SectionRecord]) -> int:
        counts = [
            count
            for section in sections
            if (count := self._metadata_int(section.metadata_json, "refined_window_count")) is not None
        ]
        return max([len(sections), *counts])

    @staticmethod
    def _metadata_int(metadata: dict[str, Any], key: str) -> int | None:
        value = metadata.get(key)
        if value is None:
            return None
        try:
            return int(str(value).strip())
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _metadata_str(metadata: dict[str, Any], key: str) -> str | None:
        value = metadata.get(key)
        if value is None:
            return None
        normalized = str(value).strip()
        return normalized or None

    @staticmethod
    def _metadata_dict_list(metadata: dict[str, Any], key: str) -> list[dict[str, Any]]:
        value = metadata.get(key)
        if not isinstance(value, list):
            return []
        return [dict(item) for item in value if isinstance(item, dict)]

    def _build_asset_record(
        self,
        *,
        source: Source,
        document: Document,
        parsed_element: ParsedElement,
        saved_sections: Sequence[SectionRecord],
        section_id_by_anchor_ref: dict[str, int],
        content_storage_key: str | None,
    ) -> AssetRecord | None:
        if parsed_element.kind not in {
            "table",
            "figure",
            "image_summary",
            "caption",
            "speaker_note",
            "ocr_region",
        }:
            return None

        section_id = self._match_section_id(
            parsed_element,
            saved_sections,
            section_id_by_anchor_ref=section_id_by_anchor_ref,
        )
        content_hash = sha256(parsed_element.text.encode("utf-8")).hexdigest()
        storage_key = self._store_asset_payload(
            parsed_element,
            fallback_storage_key=content_storage_key,
        )

        metadata_json: dict[str, Any] = dict(parsed_element.metadata)
        if parsed_element.element_id:
            metadata_json["asset_anchor_ref"] = parsed_element.element_id
            metadata_json["asset_anchor"] = asset_anchor(parsed_element.element_id)
        if parsed_element.toc_path:
            metadata_json["toc_path"] = list(parsed_element.toc_path)
        table_profile = self._table_asset_profile(parsed_element)
        if table_profile is not None:
            metadata_json.update(
                {
                    "row_count": table_profile.row_count,
                    "column_count": table_profile.column_count,
                    "estimated_tokens": table_profile.estimated_tokens,
                    "table_policy": table_profile.table_policy,
                    "asset_summary_sample": table_profile.summary_sample,
                    "sample_rows": table_profile.sample_rows,
                    "schema": table_profile.schema,
                }
            )
            asset_text_preview = table_profile.preview_text
        else:
            asset_text_preview = self._asset_text_preview(parsed_element.text)
        if asset_text_preview:
            metadata_json["asset_text_preview"] = asset_text_preview

        return AssetRecord(
            doc_id=document.doc_id,
            source_id=source.source_id,
            section_id=section_id,
            asset_type=parsed_element.kind,
            element_ref=parsed_element.element_id,
            page_no=parsed_element.page_no or 1,
            bbox={} if parsed_element.bbox is None else {
                "l": parsed_element.bbox[0],
                "t": parsed_element.bbox[1],
                "r": parsed_element.bbox[2],
                "b": parsed_element.bbox[3],
            },
            caption=parsed_element.metadata.get("caption") if parsed_element.metadata else None,
            raw_locator={},
            neighbor_section_id=section_id,
            sheet_name=self._metadata_str(metadata_json, "sheet_name"),
            row_count=self._metadata_int(metadata_json, "row_count"),
            column_count=self._metadata_int(metadata_json, "column_count"),
            sample_rows=self._metadata_dict_list(metadata_json, "sample_rows"),
            table_schema=self._metadata_dict_list(metadata_json, "schema"),
            content_hash=content_hash,
            storage_key=storage_key,
            metadata_json=metadata_json,
        )

    def _build_layout_meta_cache_record(
        self,
        *,
        source: Source,
        document: Document,
        parsed_doc: ParsedDocument,
        content_hash: str,
        object_key: str | None,
    ) -> LayoutMetaCacheRecord | None:
        layout_elements = self._serialize_layout_elements(parsed_doc.elements)
        if not layout_elements and parsed_doc.page_count is None:
            return None
        return LayoutMetaCacheRecord(
            source_id=source.source_id,
            doc_id=document.doc_id,
            content_hash=content_hash,
            object_key=object_key,
            layout_json={"elements": layout_elements},
            page_count=parsed_doc.page_count,
        )

    @staticmethod
    def _serialize_layout_elements(elements: Sequence[ParsedElement]) -> list[dict[str, object]]:
        serialized: list[dict[str, object]] = []
        for element in elements:
            record: dict[str, object] = {
                "element_id": element.element_id,
                "kind": element.kind,
                "toc_path": list(element.toc_path),
                "page_no": element.page_no,
            }
            if element.bbox is not None:
                record["bbox"] = [float(value) for value in element.bbox]
            if element.heading_level is not None:
                record["heading_level"] = element.heading_level
            if element.parent_ref is not None:
                record["parent_ref"] = element.parent_ref
            serialized.append(record)
        return serialized

    def _match_section_id(
        self,
        parsed_element: ParsedElement,
        saved_sections: Sequence[SectionRecord],
        *,
        section_id_by_anchor_ref: dict[str, int],
    ) -> int | None:
        if parsed_element.element_id:
            anchored_section_id = section_id_by_anchor_ref.get(parsed_element.element_id)
            if anchored_section_id is not None:
                return anchored_section_id

        element_path = tuple(parsed_element.toc_path)
        element_page = parsed_element.page_no

        for section in saved_sections:
            if tuple(section.toc_path) == element_path:
                return section.section_id

        if element_page is not None:
            for section in saved_sections:
                if section.page_start is not None and section.page_end is not None:
                    if section.page_start <= element_page <= section.page_end:
                        return section.section_id

        return None

    @staticmethod
    def _asset_kinds_by_ref(elements: Sequence[ParsedElement]) -> dict[str, str]:
        kinds: dict[str, str] = {}
        for element in elements:
            if element.element_id:
                kinds[element.element_id] = element.kind
        return kinds

    @staticmethod
    def _section_ids_by_anchor_ref(
        *,
        parsed_sections: Sequence[ParsedSection],
        saved_sections: Sequence[SectionRecord],
    ) -> dict[str, int]:
        section_ids: dict[str, int] = {}
        for parsed_section, saved_section in zip(parsed_sections, saved_sections, strict=False):
            for anchor_ref in iter_asset_anchor_refs(parsed_section.text):
                section_ids.setdefault(anchor_ref, saved_section.section_id)
        return section_ids

    def _store_asset_payload(
        self,
        parsed_element: ParsedElement,
        *,
        fallback_storage_key: str | None,
    ) -> str:
        if parsed_element.kind != "table":
            asset_text = self._normalize_asset_payload(parsed_element.text)
            if asset_text and self._object_store is not None:
                return self._object_store.put_bytes(
                    asset_text.encode("utf-8"),
                    suffix=f".{parsed_element.kind}.txt",
                )
            return fallback_storage_key or ""

        metadata = parsed_element.metadata or {}
        source_type = str(metadata.get("source_type", "") or "").strip()
        if source_type == SourceType.XLSX.value and fallback_storage_key:
            parquet_bytes = self._excel_source_to_parquet(
                fallback_storage_key,
                sheet_name=self._metadata_str(metadata, "sheet_name"),
            )
            if parquet_bytes is not None and self._object_store is not None:
                return self._object_store.put_bytes(parquet_bytes, suffix=".parquet")

        parquet_bytes = self._build_parquet_from_metadata(metadata, parsed_element.text)
        if parquet_bytes is not None and self._object_store is not None:
            return self._object_store.put_bytes(parquet_bytes, suffix=".parquet")
        return fallback_storage_key or ""

    def _excel_source_to_parquet(self, storage_key: str, *, sheet_name: str | None = None) -> bytes | None:
        """用 header=None 重读 Excel 源文件，走 header detector 获取干净列名，存全量 Parquet。"""
        try:
            import pandas as pd
            from openpyxl import load_workbook

            from rag.ingest.excel_formula_evaluator import ExcelFormulaEvaluator
            from rag.ingest.header_detector import HeaderKind, detect_header

            local_path = self._resolve_local_path(storage_key)
            if local_path is None:
                return None

            value_workbook = load_workbook(local_path, read_only=True, data_only=True)
            formula_workbook = load_workbook(local_path, read_only=True, data_only=False)
            try:
                selected_sheet = (
                    sheet_name
                    if sheet_name and sheet_name in formula_workbook.sheetnames
                    else formula_workbook.sheetnames[0]
                )
                worksheet = formula_workbook[selected_sheet]
                value_worksheet = value_workbook[selected_sheet]
                active_columns = self._active_excel_columns(value_worksheet, worksheet)
                if not active_columns:
                    return None
                evaluator = ExcelFormulaEvaluator(
                    value_workbook=value_workbook,
                    formula_workbook=formula_workbook,
                )
                raw = pd.DataFrame(
                    [
                        [
                            self._excel_cell_value_for_dataframe(
                                evaluator.cell_value(selected_sheet, row_number, column_number)
                            )
                            for column_number in active_columns
                        ]
                        for row_number in range(1, int(worksheet.max_row or 0) + 1)
                    ]
                )
            finally:
                formula_workbook.close()
                value_workbook.close()

            raw = raw.dropna(how="all").dropna(axis=1, how="all").fillna("")
            result = detect_header(raw)

            if result.confidence >= 0.5 and result.header_kind != HeaderKind.NONE:
                columns = result.normalized_columns
                data_start = result.data_start_row
                df = raw.iloc[data_start:].copy()
                df.columns = columns
            else:
                df = raw.copy()
                df.columns = [f"column_{index + 1}" for index in range(raw.shape[1])]

            df = df.dropna(how="all").dropna(axis=1, how="all")
            if df.empty:
                return None

            con = _duckdb.connect(":memory:")
            try:
                con.register("sheet", df)
                fd, path = tempfile.mkstemp(suffix=".parquet")
                _os_module.close(fd)
                con.execute(f"COPY sheet TO '{path}' (FORMAT PARQUET)")
                data = Path(path).read_bytes()
                _os_module.unlink(path)
                return data
            finally:
                con.close()
        except Exception:
            return None

    @classmethod
    def _active_excel_columns(cls, *worksheets: Any) -> list[int]:
        active_columns: set[int] = set()
        for worksheet in worksheets:
            max_row = int(getattr(worksheet, "max_row", 0) or 0)
            max_column = int(getattr(worksheet, "max_column", 0) or 0)
            if max_row <= 0 or max_column <= 0:
                continue
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                max_col=max_column,
                values_only=True,
            ):
                for column_index, value in enumerate(row, start=1):
                    if cls._excel_cell_value_for_dataframe(value) is not None:
                        active_columns.add(column_index)
        return sorted(active_columns)

    @staticmethod
    def _excel_cell_value_for_dataframe(value: object) -> object:
        if value is None:
            return None
        if isinstance(value, str) and not value.strip():
            return None
        return value

    def _resolve_local_path(self, storage_key: str) -> str | None:
        """从 storage_key 解析本地文件路径。"""
        if self._object_store is None:
            return None
        path_getter = getattr(self._object_store, "path_for_key", None)
        if callable(path_getter):
            try:
                p = path_getter(storage_key)
                return str(p) if p is not None and Path(p).exists() else None
            except Exception:
                pass
        return None

    @staticmethod
    def _build_parquet_from_metadata(
        metadata: dict[str, Any],
        markdown_fallback: str,
    ) -> bytes | None:
        """从 metadata 中的 columns + sample_rows 或 markdown 构建 Parquet。"""
        try:
            import pandas as pd

            columns = metadata.get("schema")
            sample_rows = metadata.get("sample_rows")
            if isinstance(columns, list) and isinstance(sample_rows, list) and columns and sample_rows:
                col_names = [c.get("name", f"col_{i}") for i, c in enumerate(columns)]
                df = pd.DataFrame(sample_rows, columns=col_names)
            else:
                df = _parse_markdown_table_to_dataframe(markdown_fallback)
                if df is None or df.empty:
                    return None

            if df.empty:
                return None

            con = _duckdb.connect(":memory:")
            try:
                con.register("sheet", df)
                fd, path = tempfile.mkstemp(suffix=".parquet")
                _os_module.close(fd)
                con.execute(f"COPY sheet TO '{path}' (FORMAT PARQUET)")
                data = Path(path).read_bytes()
                _os_module.unlink(path)
                return data
            finally:
                con.close()
        except Exception:
            return None

    @staticmethod
    def _normalize_asset_payload(text: str) -> str:
        return text.replace("\r\n", "\n").replace("\r", "\n").strip()

    def _asset_text_preview(self, text: str, *, max_tokens: int = 1000) -> str:
        normalized = self._normalize_asset_payload(text)
        if not normalized:
            return ""
        try:
            return self._section_refiner.token_accounting.clip(normalized, max_tokens).strip()
        except Exception:
            return normalized

    def _table_asset_profile(self, parsed_element: ParsedElement) -> TableAssetProfile | None:
        if parsed_element.kind != "table":
            return None
        metadata_profile = self._table_asset_profile_from_metadata(parsed_element.metadata)
        if metadata_profile is not None:
            return metadata_profile
        return profile_markdown_table(
            parsed_element.text,
            token_accounting=self._section_refiner.token_accounting,
        )

    def _table_asset_profile_from_metadata(self, metadata: dict[str, Any]) -> TableAssetProfile | None:
        table_policy = str(metadata.get("table_policy", "") or "").strip()
        row_count = self._metadata_int(metadata, "row_count")
        column_count = self._metadata_int(metadata, "column_count")
        estimated_tokens = self._metadata_int(metadata, "estimated_tokens")
        summary_sample = metadata.get("asset_summary_sample")
        if (
            not table_policy
            or row_count is None
            or column_count is None
            or estimated_tokens is None
            or not isinstance(summary_sample, str)
            or not summary_sample.strip()
        ):
            return None
        preview = metadata.get("asset_text_preview")
        return TableAssetProfile(
            row_count=row_count,
            column_count=column_count,
            estimated_tokens=estimated_tokens,
            table_policy=table_policy,
            summary_sample=summary_sample.strip(),
            preview_text=preview.strip() if isinstance(preview, str) and preview.strip() else summary_sample.strip(),
            columns=[
                str(item.get("name", "")).strip()
                for item in self._metadata_dict_list(metadata, "schema")
                if str(item.get("name", "")).strip()
            ],
            schema=self._metadata_dict_list(metadata, "schema"),
            sample_rows=self._metadata_dict_list(metadata, "sample_rows"),
        )

    # ============================================================
    # L2 summary + vector writes
    # ============================================================

    def _summary_records_for_document(
        self,
        *,
        source: Source,
        document: Document,
        parsed_doc: ParsedDocument,
        saved_sections: Sequence[SectionRecord],
        saved_assets: Sequence[AssetRecord],
    ) -> list[DocSummaryRecord | SectionSummaryRecord | AssetSummaryRecord]:
        section_records: list[SectionSummaryRecord] = []
        for parsed_section, section_record in zip(parsed_doc.sections, saved_sections, strict=False):
            section_records.append(
                self._section_summary_record(
                    source=source,
                    document=document,
                    section_record=section_record,
                    parsed_section=parsed_section,
                )
            )
        asset_records: list[AssetSummaryRecord] = []
        for asset_record in saved_assets:
            asset_records.append(
                self._asset_summary_record(
                    source=source,
                    document=document,
                    asset_record=asset_record,
                )
            )
        doc_record = self._doc_summary_record(
            source=source,
            document=document,
            parsed_doc=parsed_doc,
            section_summary_records=section_records,
            asset_summary_records=asset_records,
        )
        return [doc_record, *section_records, *asset_records]

    def _write_prepared_summary_records(self, prepared_items: Sequence[_PreparedIngestItem]) -> None:
        records_to_embed: list[DocSummaryRecord | SectionSummaryRecord | AssetSummaryRecord] = []
        vector_indexes: list[tuple[DocSummaryRecord | SectionSummaryRecord | AssetSummaryRecord, int]] = []

        for prepared in prepared_items:
            records = self._summary_records_for_document(
                source=prepared.source,
                document=prepared.document,
                parsed_doc=prepared.parsed_doc,
                saved_sections=prepared.saved_sections,
                saved_assets=prepared.saved_assets,
            )
            for record in records:
                vector_index = len(records_to_embed)
                records_to_embed.append(record)
                vector_indexes.append((record, vector_index))

        if not records_to_embed:
            return

        vectors = self._embedder.embed([record.summary_text for record in records_to_embed])
        if len(vectors) != len(records_to_embed):
            raise RuntimeError(
                f"embedding count mismatch: expected {len(records_to_embed)}, got {len(vectors)}"
            )

        upsert_items = [(record, vectors[vector_index]) for record, vector_index in vector_indexes]
        upsert_records = getattr(self._summary_repo, "upsert_records", None)
        if callable(upsert_records):
            upsert_records(upsert_items, embedding_space=self._embedding_space)
            return
        for record, vector in upsert_items:
            self._summary_repo.upsert_record(record, vector, embedding_space=self._embedding_space)

    def _doc_summary_record(
        self,
        *,
        source: Source,
        document: Document,
        parsed_doc: ParsedDocument,
        section_summary_records: Sequence[SectionSummaryRecord],
        asset_summary_records: Sequence[AssetSummaryRecord],
    ) -> DocSummaryRecord:
        summarize_doc = getattr(self._summarizer, "summarize_doc_with_metadata", None)
        if callable(summarize_doc):
            summary = summarize_doc(
                document_title=document.title or parsed_doc.title or "document",
                section_summaries=[record.summary_text for record in section_summary_records],
                asset_summaries=[record.summary_text for record in asset_summary_records],
            )
        else:
            summary = RetrievalSummaryResult(
                text=self._fallback_doc_summary(
                    document_title=document.title or parsed_doc.title or "document",
                    child_summaries=[
                        *(record.summary_text for record in section_summary_records),
                        *(record.summary_text for record in asset_summary_records),
                    ],
                ),
                method="fallback_doc_reduce",
                provider_name=None,
                model_name=None,
            )

        return DocSummaryRecord(
            doc_id=document.doc_id,
            source_id=source.source_id,
            version_group_id=document.version_group_id,
            version_no=document.version_no,
            doc_status=document.doc_status,
            effective_date=document.effective_date,
            is_active=document.is_active,
            index_ready=True,
            tenant_id=document.tenant_id,
            department_id=document.department_id,
            auth_tag=document.auth_tag,
            source_type=source.source_type,
            embedding_model_id=document.embedding_model_id,
            partition_key=PartitionKey.HOT,
            title=document.title,
            summary_text=summary.text,
            metadata_json={
                "summary_generation": {
                    "method": summary.method,
                    "provider_name": summary.provider_name,
                    "model_name": summary.model_name,
                    "fallback_reason": summary.fallback_reason,
                }
            },
        )

    def _section_summary_record(
        self,
        *,
        source: Source,
        document: Document,
        section_record: SectionRecord,
        parsed_section: ParsedSection,
    ) -> SectionSummaryRecord:
        summary = self._summarizer.summarize_section_with_metadata(
            parsed_section,
            document.title or "document",
        )

        return SectionSummaryRecord(
            section_id=section_record.section_id,
            doc_id=document.doc_id,
            source_id=source.source_id,
            version_group_id=document.version_group_id,
            version_no=document.version_no,
            doc_status=document.doc_status,
            effective_date=document.effective_date,
            is_active=document.is_active,
            index_ready=True,
            tenant_id=document.tenant_id,
            department_id=document.department_id,
            auth_tag=document.auth_tag,
            source_type=source.source_type,
            embedding_model_id=document.embedding_model_id,
            partition_key=PartitionKey.HOT,
            page_start=section_record.page_start,
            page_end=section_record.page_end,
            section_kind=section_record.section_kind,
            toc_path=list(section_record.toc_path),
            summary_text=summary.text,
            metadata_json={
                "summary_generation": {
                    "method": summary.method,
                    "provider_name": summary.provider_name,
                    "model_name": summary.model_name,
                    "fallback_reason": summary.fallback_reason,
                }
            },
        )

    def _asset_summary_record(
        self,
        *,
        source: Source,
        document: Document,
        asset_record: AssetRecord,
    ) -> AssetSummaryRecord:
        asset_text = self._read_asset_text_for_summary(asset_record)
        toc_path = asset_record.metadata_json.get("toc_path", [])
        if not isinstance(toc_path, list):
            toc_path = []
        summary = self._summarizer.summarize_asset_with_metadata(
            asset_type=asset_record.asset_type,
            asset_text=asset_text,
            document_title=document.title or "document",
            toc_path=[str(part) for part in toc_path],
            caption=asset_record.caption,
        )

        return AssetSummaryRecord(
            asset_id=asset_record.asset_id,
            doc_id=document.doc_id,
            source_id=source.source_id,
            section_id=asset_record.section_id,
            version_group_id=document.version_group_id,
            version_no=document.version_no,
            doc_status=document.doc_status,
            effective_date=document.effective_date,
            is_active=document.is_active,
            index_ready=True,
            tenant_id=document.tenant_id,
            department_id=document.department_id,
            auth_tag=document.auth_tag,
            embedding_model_id=document.embedding_model_id,
            partition_key=PartitionKey.HOT,
            asset_type=asset_record.asset_type,
            page_no=asset_record.page_no,
            caption=asset_record.caption,
            summary_text=summary.text,
            metadata_json={
                "summary_generation": {
                    "method": summary.method,
                    "provider_name": summary.provider_name,
                    "model_name": summary.model_name,
                    "fallback_reason": summary.fallback_reason,
                }
            },
        )

    def _read_asset_text_for_summary(self, asset_record: AssetRecord) -> str:
        sample = asset_record.metadata_json.get("asset_summary_sample")
        if isinstance(sample, str) and sample.strip():
            return sample.strip()
        preview = asset_record.metadata_json.get("asset_text_preview")
        if isinstance(preview, str) and preview.strip():
            return preview.strip()
        if self._object_store is None or not asset_record.storage_key:
            return asset_record.caption or asset_record.asset_type
        try:
            reader = getattr(self._object_store, "read_byte_range", None)
            if callable(reader):
                raw = reader(asset_record.storage_key, 0, 24_000)
            else:
                raw = self._object_store.read_bytes(asset_record.storage_key)[:24_000]
        except Exception:
            return asset_record.caption or asset_record.asset_type
        text = raw.decode("utf-8", errors="ignore").strip()
        return self._asset_text_preview(text) or asset_record.caption or asset_record.asset_type

    # ============================================================
    # helpers
    # ============================================================

    def _resolve_raw_bytes(self, request: IngestRequest) -> bytes:
        if request.content_text is not None:
            return request.content_text.encode("utf-8")
        if request.raw_bytes is not None:
            return request.raw_bytes
        if request.file_path is not None:
            return request.file_path.read_bytes()

        fallback_path = Path(request.location)
        if fallback_path.exists():
            return fallback_path.read_bytes()

        raise ValueError(f"No content available for: {request.location}")

    def _fallback_doc_summary(self, *, document_title: str, child_summaries: Sequence[str]) -> str:
        text = normalize_whitespace(" ".join([document_title, *child_summaries]))
        if not text:
            return normalize_whitespace(document_title)
        try:
            return self._section_refiner.token_accounting.clip(text, 220).strip()
        except Exception:
            return text

    def _save_processing_state(
        self,
        *,
        doc_id: int,
        source_id: int,
        stage: str,
        status: str,
        error_message: str | None = None,
    ) -> None:
        record = ProcessingStateRecord(
            doc_id=doc_id,
            source_id=source_id,
            stage=stage,
            status=status,
            attempts=1,
            priority="normal",
            worker_id=None,
            lease_expires_at=None,
            error_message=error_message,
            metadata_json={},
        )
        self._metadata_repo.save_processing_state(record)
