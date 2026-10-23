"""Source-preserving tabular inspection, with no implicit header inference."""
from __future__ import annotations

import hashlib
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from agent_runtime.workspace import open_workspace
from tests.agent.test_builtin_data_tools import _execute, _tools


async def inspect(tmp_path: Path, name: str, **options):
    workspace = open_workspace(tmp_path)
    result = await _execute(_tools(workspace)["inspect_data_file"], {"path": name, **options}, workspace=workspace)
    assert not result.result.is_error, result.result.error_message
    return result.result.structured_content


@pytest.mark.anyio
async def test_title_units_multilevel_headers_and_blank_rows_are_raw_cells(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Quarterly report"])
    ws.merge_cells("A1:C1")
    ws.append(["Unit: USD"])
    ws.append([None, "Previous", "Current"])
    ws.append(["Product", "Revenue", "Revenue"])
    ws.append(["A", 10, 20])
    ws.append([])
    ws.append(["Notes", "provisional"])
    path = tmp_path / "sales.xlsx"
    wb.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    out = await inspect(tmp_path, "sales.xlsx")
    grid = out["tables"][0]
    assert "headers" not in grid
    assert grid["rows"] == (
        ("Quarterly report", "", ""), ("Unit: USD", "", ""),
        ("", "Previous", "Current"), ("Product", "Revenue", "Revenue"),
        ("A", "10", "20"), ("", "", ""), ("Notes", "provisional", ""),
    )
    assert grid["row_numbers"] == tuple(range(1, 8))
    assert grid["column_numbers"] == (1, 2, 3)
    assert not grid["truncated"]
    assert out["sha256"] == before == hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.mark.anyio
@pytest.mark.parametrize("extension,separator", [("csv", ","), ("tsv", "\t")])
async def test_headerless_delimited_records_keep_first_record_and_blank_records(tmp_path, extension, separator):
    (tmp_path / f"raw.{extension}").write_text(
        f'001{separator}NA\n\n002{separator}"two\nlines"\n', encoding="utf-8-sig",
    )
    out = await inspect(tmp_path, f"raw.{extension}")
    grid = out["tables"][0]
    assert "headers" not in grid
    assert grid["rows"] == (("001", "NA"), ("", ""), ("002", "two\nlines"))
    assert grid["row_numbers"] == (1, 2, 3)  # Logical CSV records, not physical lines.
    assert grid["row_count"] == 3


@pytest.mark.anyio
@pytest.mark.parametrize("extension", ["xlsx", "csv"])
async def test_row_and_column_pages_retain_source_coordinates(tmp_path, extension):
    rows = [[f"r{r}c{c}" for c in range(1, 6)] for r in range(1, 6)]
    if extension == "xlsx":
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(tmp_path / "grid.xlsx")
    else:
        (tmp_path / "grid.csv").write_text("\n".join(",".join(r) for r in rows))
    out = await inspect(tmp_path, f"grid.{extension}", start_row=2, start_column=2, max_rows=2, max_columns=2)
    grid = out["tables"][0]
    assert grid["rows"] == (("r2c2", "r2c3"), ("r3c2", "r3c3"))
    assert grid["row_numbers"] == (2, 3)
    assert grid["column_numbers"] == (2, 3)
    assert grid["next_row"] == 4
    assert grid["next_column"] == 4
    assert grid["truncated"] and out["truncated"]
    tail = (await inspect(tmp_path, f"grid.{extension}", start_row=4, start_column=4,
                          max_rows=2, max_columns=2))["tables"][0]
    assert tail["rows"] == (("r4c4", "r4c5"), ("r5c4", "r5c5"))
    assert tail["next_row"] is None and tail["next_column"] is None
    assert tail["row_count"] == 5 and tail["column_count"] == 5


@pytest.mark.anyio
async def test_formulas_errors_and_truncated_cell_text_are_identified(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["=SUM(1,2)", "#DIV/0!", "x" * 2500])
    wb.save(tmp_path / "formula.xlsx")
    grid = (await inspect(tmp_path, "formula.xlsx"))["tables"][0]
    assert grid["rows"][0][0] == "=SUM(1,2)"  # No cached value exists; formula must not disappear.
    assert grid["cell_types"][0][:2] == ("f", "e")
    assert grid["truncated_cells"] == ((1, 3),)
    assert grid["truncated"]


@pytest.mark.anyio
async def test_sheet_selection_and_sheet_paging(tmp_path):
    wb = Workbook()
    wb.active.title = "First"
    wb.active.append(["first"])
    wb.create_sheet("Second").append(["second"])
    wb.create_sheet("Third").append(["third"])
    wb.save(tmp_path / "sheets.xlsx")
    first = await inspect(tmp_path, "sheets.xlsx", max_sheets=1)
    assert first["next_sheet_offset"] == 1 and first["truncated"]
    second = await inspect(tmp_path, "sheets.xlsx", sheet_offset=1, max_sheets=1)
    assert second["tables"][0]["name"] == "Second" and second["next_sheet_offset"] == 2
    selected = await inspect(tmp_path, "sheets.xlsx", sheet_name="Third")
    assert [t["name"] for t in selected["tables"]] == ["Third"]
    assert not selected["truncated"] and selected["next_sheet_offset"] is None


@pytest.mark.anyio
async def test_declared_excel_table_does_not_discard_surrounding_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Report title"])
    ws.append(["Item", "Value"])
    ws.append(["A", 3])
    ws.add_table(Table(displayName="SalesTable", ref="A2:B3"))
    ws.append(["Unrelated note"])
    wb.save(tmp_path / "table.xlsx")
    grid = (await inspect(tmp_path, "table.xlsx"))["tables"][0]
    assert grid["rows"] == (("Report title", ""), ("Item", "Value"), ("A", "3"), ("Unrelated note", ""))
    assert "headers" not in grid


@pytest.mark.anyio
async def test_streaming_preview_recovers_cells_outside_incorrect_declared_dimensions(tmp_path):
    wb = Workbook()
    wb.active.append(["title"])
    wb.active["C4"] = "must remain visible"
    path = tmp_path / "dimensions.xlsx"
    wb.save(path)
    with ZipFile(path) as archive:
        entries = {n: archive.read(n) for n in archive.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b'ref="A1:C4"', b'ref="A1:A1"')
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    grid = (await inspect(tmp_path, "dimensions.xlsx"))["tables"][0]
    assert grid["rows"][3][2] == "must remain visible"
    assert grid["row_count"] == 4 and grid["column_count"] == 3


@pytest.mark.anyio
async def test_invalid_encoding_fails_instead_of_replacing_source_bytes(tmp_path):
    (tmp_path / "bad.csv").write_bytes(b"name,value\nA,\xff\n")
    workspace = open_workspace(tmp_path)
    result = await _execute(_tools(workspace)["inspect_data_file"], {"path": "bad.csv"}, workspace=workspace)
    assert result.result.is_error
    assert result.result.error_code == "data_inspection_failed"


@pytest.mark.anyio
async def test_empty_and_past_end_grids_are_not_content_evidence(tmp_path):
    wb = Workbook()
    wb.save(tmp_path / "empty.xlsx")
    empty = (await inspect(tmp_path, "empty.xlsx"))["tables"][0]
    assert empty["rows"] == () and empty["row_count"] == 0
    (tmp_path / "one.csv").write_text("a,b\n")
    past = (await inspect(tmp_path, "one.csv", start_row=10))["tables"][0]
    assert past["rows"] == () and past["row_numbers"] == ()
    assert past["row_count"] == 1 and past["next_row"] is None


@pytest.mark.parametrize("entrypoint", ["sdk", "cli"])
def test_public_entrypoints_preserve_raw_grid_in_model_request(tmp_path, monkeypatch, entrypoint):
    import asyncio
    import json
    from dataclasses import replace

    from typer.testing import CliRunner

    from agent_runtime import Agent
    from agent_runtime import cli as cli_module
    from agent_runtime.harness import HarnessModelResponse, HarnessToolCall, RolloutStore
    from tests.agent.harness.test_subagent_budget import SmallAnswerModel

    class InspectModel(SmallAnswerModel):
        def __init__(self):
            self.grid = None

        def prepare(self, request):
            prepared = super().prepare(request)
            return replace(prepared, dispatch_payload=request, request_ref={
                **prepared.request_ref,
                "toolset_revision": prepared.tool_hash,
                "exposed_tool_names": [tool.definition.name for tool in request.tools],
            })

        async def dispatch(self, prepared, *, delta_sink=None):
            messages = prepared.dispatch_payload.messages
            results = [m for m in messages if m.role == "tool"]
            if not results:
                return HarnessModelResponse(text="", provider_response_id="inspect", usage={
                    "input_tokens": 2, "output_tokens": 3, "total_tokens": 5,
                }, tool_calls=(HarnessToolCall(id="inspect-call", name="inspect_data_file",
                    arguments={"path": "source.xlsx"}),))
            content = json.loads(results[-1].content)
            self.grid = content["structured_content"]["tables"][0]
            assert self.grid["rows"] == [["Report", ""], ["Unit: USD", ""], ["Item", "Value"], ["A", "7"]]
            assert "headers" not in self.grid
            assert self.grid["row_numbers"] == [1, 2, 3, 4]
            return HarnessModelResponse(text="Observed raw source cells.", provider_response_id="answer", usage={
                "input_tokens": 2, "output_tokens": 3, "total_tokens": 5,
            })

    workspace_path = tmp_path / "workspace"
    workspace_path.mkdir()
    monkeypatch.chdir(workspace_path)
    monkeypatch.setenv("PRAXIS_CONFIG_DIR", str(tmp_path / "host-config"))
    monkeypatch.setenv("PRAXIS_MODEL_REGISTRY_PATH", str(tmp_path / "host-config/models.yaml"))
    wb = Workbook()
    for row in [["Report"], ["Unit: USD"], ["Item", "Value"], ["A", 7]]:
        wb.active.append(row)
    wb.save(workspace_path / "source.xlsx")
    database = tmp_path / "rollout.sqlite"
    model = InspectModel()
    if entrypoint == "sdk":
        agent = Agent(workspace_path=workspace_path, checkpoint_db=database, model_session_path=None,
                      enable_workspace_mcp=False)
        agent._model_factory = lambda: model
        result = asyncio.run(agent.run("Inspect source.xlsx", max_steps=3, max_tokens_total=1000))
        assert result.status == "done"
    else:
        original = cli_module._create_agent_facade

        def facade(**options):
            agent = original(**options)
            agent._model_factory = lambda: model
            return agent

        monkeypatch.setattr(cli_module, "_create_agent_facade", facade)
        result = CliRunner().invoke(cli_module.agent_app, [
            "run", "Inspect source.xlsx", "--rollout-db", str(database), "--non-interactive",
            "--no-web-tools", "--disable-workspace-mcp", "--max-steps", "3", "--max-tokens-total", "1000",
        ])
        assert result.exit_code == 0, result.output
    assert model.grid is not None
    with RolloutStore(database) as store:
        assert store.verify().valid


@pytest.mark.anyio
async def test_array_formula_expression_is_preserved(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula

    wb = Workbook()
    wb.active["A1"] = ArrayFormula(ref="A1:A3", text="=ROW(A1:A3)")
    wb.save(tmp_path / "array.xlsx")
    grid = (await inspect(tmp_path, "array.xlsx"))["tables"][0]
    assert grid["rows"][0][0] == "=ROW(A1:A3)"
    assert grid["cell_types"][0][0] == "f"


@pytest.mark.anyio
async def test_large_unicode_grid_retains_coordinates_and_truncation_feedback(tmp_path):
    (tmp_path / "long.csv").write_text("\n".join([",".join(["字" * 2000] * 20)] * 10))
    out = await inspect(tmp_path, "long.csv")
    grid = out["tables"][0]
    assert grid["row_numbers"] == tuple(range(1, 11))
    assert grid["column_numbers"] == tuple(range(1, 21))
    assert grid["truncated"] and grid["truncated_cells"]


@pytest.mark.anyio
async def test_formats_remain_distinct_from_stored_values(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 0.15
    ws["A1"].number_format = "0%"
    ws["B1"] = 12
    ws["B1"].number_format = "0000"
    ws["C1"] = "0012"
    wb.save(tmp_path / "formats.xlsx")
    grid = (await inspect(tmp_path, "formats.xlsx"))["tables"][0]
    assert grid["rows"][0] == ("0.15", "12", "0012")
    assert grid["cell_types"][0] == ("n", "n", "s")
    assert grid["number_formats"][0] == ("0%", "0000", "")


@pytest.mark.anyio
async def test_data_table_formula_metadata_is_deterministic(tmp_path):
    import json

    from openpyxl.worksheet.formula import DataTableFormula

    wb = Workbook()
    wb.active["A1"] = DataTableFormula(ref="A1:B3", r1="C1")
    wb.save(tmp_path / "data-table.xlsx")
    grid = (await inspect(tmp_path, "data-table.xlsx"))["tables"][0]
    formula = json.loads(grid["rows"][0][0])
    assert formula["formula_type"] == "data_table"
    assert formula["ref"] == "A1:B3" and formula["r1"] == "C1"
    assert grid["cell_types"][0][0] == "f"


@pytest.mark.anyio
async def test_cell_budget_preserves_remaining_row_and_sheet_pages(tmp_path):
    wb = Workbook()
    for ws in [wb.active, wb.create_sheet("Second")]:
        for _ in range(25):
            ws.append(["v"] * 100)
    wb.save(tmp_path / "wide.xlsx")
    out = await inspect(tmp_path, "wide.xlsx", max_rows=100, max_columns=100)
    assert sum(len(t["rows"]) * len(t["column_numbers"]) for t in out["tables"]) <= 2000
    assert out["tables"][0]["next_row"] == 21
    assert out["next_sheet_offset"] == 1 and out["truncated"]
    tail = await inspect(tmp_path, "wide.xlsx", sheet_name=wb.active.title, start_row=21,
                         max_rows=100, max_columns=100)
    assert tail["tables"][0]["row_numbers"] == (21, 22, 23, 24, 25)
    assert tail["tables"][0]["next_row"] is None


@pytest.mark.anyio
async def test_old_inspection_execution_revision_cannot_be_restored(tmp_path):
    from agent_runtime.harness import Session
    from tests.agent.harness.test_subagent_budget import SmallAnswerModel

    workspace = open_workspace(tmp_path)
    tool = _tools(workspace)["inspect_data_file"]
    async with await Session.open(database=tmp_path / "history.sqlite", workspace=tmp_path,
                                  model=SmallAnswerModel(), tools={"inspect_data_file": tool}) as runtime:
        orchestrator = runtime.tool_orchestrator
        with pytest.raises(RuntimeError, match="original operation tool is unavailable: inspect_data_file"):
            orchestrator.restore_tools({"inspect_data_file": "builtin-inspect-data-file-v1"})
        assert orchestrator.restore_tools({"inspect_data_file": tool.execution_revision}) == (tool,)


@pytest.mark.anyio
async def test_exhausted_text_budget_does_not_add_unbudgeted_markers(tmp_path):
    (tmp_path / "budget.csv").write_text("\n".join([",".join(["x"] * 100)] * 20))
    out = await inspect(tmp_path, "budget.csv", max_rows=100, max_columns=100, max_characters=100)
    grid = out["tables"][0]
    assert sum(len(s) for row in grid["rows"] for s in row) <= 100
    assert grid["truncated_cells"] and grid["truncated"]
